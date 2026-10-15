"""Extração de texto dos materiais e seleção dos trechos mais relevantes.

A seleção usa BM25 (algoritmo clássico de busca) em Python puro: dispensa banco
vetorial, não baixa modelos e não envia o material inteiro para a IA sem necessidade.
"""
import io
import math
import re
import unicodedata
from collections import Counter

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

FORMATOS = ["pdf", "docx", "pptx"]

_STOPWORDS = set("""
a o e é de da do das dos em no na nos nas um uma uns umas para por com sem que se ao aos
às as os ou como mais mas não sua seu suas seus ser são pelo pela pelos pelas este esta
isso isto esse essa entre sobre quando onde qual quais deve devem pode podem também já
muito muita está estão foi foram ter tem têm há cada após antes até
""".split())


def extrair_texto(nome_arquivo: str, dados: bytes) -> str:
    """Extrai o texto de PDF, DOCX (incluindo tabelas) ou PPTX (incluindo notas)."""
    extensao = nome_arquivo.rsplit(".", 1)[-1].lower()
    buffer = io.BytesIO(dados)
    partes: list[str] = []

    if extensao == "pdf":
        leitor = PdfReader(buffer)
        if leitor.is_encrypted and not leitor.decrypt(""):
            raise ValueError("O PDF está protegido por senha.")
        partes += [pagina.extract_text() or "" for pagina in leitor.pages]

    elif extensao == "docx":
        documento = Document(buffer)
        partes += [p.text for p in documento.paragraphs]
        for tabela in documento.tables:
            for linha in tabela.rows:
                partes.append(" | ".join(c.text.strip() for c in linha.cells))

    elif extensao == "pptx":
        apresentacao = Presentation(buffer)
        for numero, slide in enumerate(apresentacao.slides, 1):
            partes.append(f"[Slide {numero}]")
            for forma in slide.shapes:
                if forma.has_text_frame:
                    partes += [p.text for p in forma.text_frame.paragraphs]
            if slide.has_notes_slide:
                notas = slide.notes_slide.notes_text_frame
                if notas is not None and notas.text.strip():
                    partes.append("Notas do apresentador: " + notas.text)
    else:
        raise ValueError("Formato não suportado. Use PDF, DOCX ou PPTX.")

    return "\n".join(p for p in partes if p and p.strip())


def dividir_em_trechos(texto: str, tamanho: int = 1200) -> list[str]:
    paragrafos = [p.strip() for p in re.split(r"\n+", texto) if p.strip()]
    trechos: list[str] = []
    atual = ""
    for paragrafo in paragrafos:
        while len(paragrafo) > tamanho:
            if atual:
                trechos.append(atual)
                atual = ""
            trechos.append(paragrafo[:tamanho])
            paragrafo = paragrafo[tamanho:]
        if atual and len(atual) + len(paragrafo) + 1 > tamanho:
            trechos.append(atual)
            atual = paragrafo
        else:
            atual = f"{atual}\n{paragrafo}" if atual else paragrafo
    if atual:
        trechos.append(atual)
    return trechos


def _tokens(texto: str) -> list[str]:
    t = unicodedata.normalize("NFKD", texto.lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    # Radical simples (5 letras): aproxima "atendimento" de "atender", "cliente" de "clientes"
    return [p[:5] for p in re.findall(r"[a-z0-9]{3,}", t) if p not in _STOPWORDS]


def selecionar_trechos(texto: str, consulta: str, k: int = 5, limite_integral: int = 10_000) -> list[str]:
    """Material curto vai inteiro; material longo é filtrado pelos k trechos mais relevantes."""
    trechos = dividir_em_trechos(texto)
    if len(texto) <= limite_integral or len(trechos) <= k:
        return trechos

    documentos = [_tokens(t) for t in trechos]
    termos = set(_tokens(consulta))
    if not termos:
        return trechos[:k]

    n = len(documentos)
    media = sum(len(d) for d in documentos) / n or 1
    frequencia_doc = Counter(termo for d in documentos for termo in set(d))
    k1, b = 1.5, 0.75

    pontuacoes = []
    for indice, documento in enumerate(documentos):
        tf = Counter(documento)
        pontos = 0.0
        for termo in termos:
            if termo not in tf:
                continue
            idf = math.log(1 + (n - frequencia_doc[termo] + 0.5) / (frequencia_doc[termo] + 0.5))
            pontos += idf * tf[termo] * (k1 + 1) / (tf[termo] + k1 * (1 - b + b * len(documento) / media))
        pontuacoes.append((pontos, indice))

    melhores = sorted(i for _, i in sorted(pontuacoes, reverse=True)[:k])
    return [trechos[i] for i in melhores]
